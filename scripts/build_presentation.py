"""Generate the NEXUS Corp CTF presentation (.pptx).

Light-mode, minimal style, accent colors taken from the web app palette.
Run:  python scripts/build_presentation.py
Output: NexusCorp_CTF_Presentation.pptx (at repo root)
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


# ── Palette (light mode, accents from nexus.css) ─────────────────────────
BG         = RGBColor(0xFF, 0xFF, 0xFF)
BG_SOFT    = RGBColor(0xF5, 0xF7, 0xFA)
TEXT       = RGBColor(0x14, 0x17, 0x25)
TEXT_DIM   = RGBColor(0x5A, 0x5D, 0x72)
ACCENT     = RGBColor(0x00, 0xA3, 0xAD)   # cyan (darker for contrast)
RED        = RGBColor(0xE5, 0x1A, 0x3D)
GREEN      = RGBColor(0x00, 0x99, 0x5C)
ORANGE     = RGBColor(0xE5, 0x7A, 0x00)
PURPLE     = RGBColor(0x8B, 0x3E, 0xE5)
BORDER     = RGBColor(0xE3, 0xE6, 0xEC)

FONT_TITLE = "Calibri"
FONT_BODY  = "Calibri"
FONT_MONO  = "Consolas"


# ── Setup ────────────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW, SH = prs.slide_width, prs.slide_height


def add_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.shadow.inherit = False
    return s


def add_text(slide, left, top, width, height, text,
             size=18, bold=False, color=TEXT, font=FONT_BODY,
             align=PP_ALIGN.LEFT, line_spacing=1.15):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def add_rect(slide, left, top, width, height, fill, line=None):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
        r.line.width = Pt(0.75)
    r.shadow.inherit = False
    return r


def add_header(slide, eyebrow, title, accent=ACCENT):
    add_rect(slide, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.45), accent)
    add_text(slide, Inches(0.85), Inches(0.50), Inches(10), Inches(0.35),
             eyebrow.upper(), size=11, bold=True, color=accent,
             font=FONT_TITLE)
    add_text(slide, Inches(0.85), Inches(0.82), Inches(12), Inches(0.7),
             title, size=30, bold=True, color=TEXT, font=FONT_TITLE)
    add_rect(slide, Inches(0.85), Inches(1.55), Inches(11.6), Inches(0.02),
             BORDER)


def add_footer(slide, page_num):
    add_text(slide, Inches(0.6), Inches(7.05), Inches(6), Inches(0.3),
             "NEXUS Corp CTF · Carlos · UP", size=9, color=TEXT_DIM)
    add_text(slide, Inches(11.5), Inches(7.05), Inches(1.5), Inches(0.3),
             str(page_num).zfill(2), size=9, color=TEXT_DIM,
             align=PP_ALIGN.RIGHT)


def add_code(slide, left, top, width, height, code, size=12):
    add_rect(slide, left, top, width, height, BG_SOFT, line=BORDER)
    tb = slide.shapes.add_textbox(
        left + Inches(0.15), top + Inches(0.12),
        width - Inches(0.3), height - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = code.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = line
        r.font.name = FONT_MONO
        r.font.size = Pt(size)
        r.font.color.rgb = TEXT


def add_bullets(slide, left, top, width, height, items,
                size=16, color=TEXT, bullet_color=None):
    if bullet_color is None:
        bullet_color = ACCENT
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.3
        p.space_after = Pt(6)
        rb = p.add_run()
        rb.text = "▸  "
        rb.font.name = FONT_BODY
        rb.font.size = Pt(size)
        rb.font.color.rgb = bullet_color
        rb.font.bold = True
        r = p.add_run()
        r.text = item
        r.font.name = FONT_BODY
        r.font.size = Pt(size)
        r.font.color.rgb = color


def add_chip(slide, left, top, text, color):
    w, h = Inches(1.3), Inches(0.32)
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    r.fill.solid()
    r.fill.fore_color.rgb = BG
    r.line.color.rgb = color
    r.line.width = Pt(1)
    r.shadow.inherit = False
    add_text(slide, left, top + Inches(0.02), w, h, text.upper(),
             size=9, bold=True, color=color, align=PP_ALIGN.CENTER,
             font=FONT_TITLE)


# ═════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title / Cover
# ═════════════════════════════════════════════════════════════════════════
s = add_slide()
add_rect(s, 0, 0, Inches(0.35), SH, ACCENT)

add_text(s, Inches(1.0), Inches(1.7), Inches(12), Inches(0.5),
         "CAPTURE THE FLAG", size=14, bold=True, color=ACCENT,
         font=FONT_TITLE)
add_text(s, Inches(1.0), Inches(2.15), Inches(12), Inches(1.5),
         "NEXUS Corp CTF", size=54, bold=True, color=TEXT,
         font=FONT_TITLE)
add_text(s, Inches(1.0), Inches(3.35), Inches(12), Inches(1.0),
         "Seis retos, seis vulnerabilidades web reales.",
         size=22, color=TEXT_DIM, font=FONT_TITLE)

add_rect(s, Inches(1.0), Inches(4.3), Inches(0.6), Inches(0.04), ACCENT)

add_text(s, Inches(1.0), Inches(4.5), Inches(12), Inches(0.4),
         "Hackeo Ético y Recuperación Ante Desastres  ·  UP 8vo Semestre",
         size=14, color=TEXT_DIM, font=FONT_TITLE)
add_text(s, Inches(1.0), Inches(4.85), Inches(12), Inches(0.4),
         "Carlos  ·  2026", size=14, bold=True, color=TEXT,
         font=FONT_TITLE)

# tech chips
chips = [("Flask", ACCENT), ("SQLite", ACCENT), ("Hydra", RED),
         ("sqlmap", RED), ("Hashcat", ORANGE), ("ffuf", PURPLE)]
for i, (t, c) in enumerate(chips):
    add_chip(s, Inches(1.0 + i * 1.45), Inches(6.2), t, c)


# ═════════════════════════════════════════════════════════════════════════
# SLIDE 2 — Introducción al proyecto
# ═════════════════════════════════════════════════════════════════════════
s = add_slide()
add_header(s, "01 · introducción", "¿Qué es NEXUS Corp CTF?")

add_text(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(0.8),
         "Plataforma web educativa (Flask + SQLite) que simula la "
         "infraestructura comprometida de una corporación ficticia, "
         "NEXUS Corp.",
         size=17, color=TEXT)

add_text(s, Inches(0.85), Inches(2.8), Inches(11.6), Inches(0.4),
         "El jugador progresa resolviendo 6 retos encadenados, cada uno "
         "réplica de una vulnerabilidad real del OWASP Top 10.",
         size=17, color=TEXT)

# 3 columns
col_w = Inches(3.75)
col_y = Inches(4.0)
col_h = Inches(2.4)
for i, (title, body, color) in enumerate([
    ("Objetivo didáctico",
     "Aprender pentesting web con herramientas estándar de la "
     "industria en un entorno controlado.",
     ACCENT),
    ("Narrativa encadenada",
     "Cada reto entrega credenciales o pistas que habilitan el "
     "siguiente. No hay atajos.",
     ORANGE),
    ("Stack minimalista",
     "Flask, SQLite, Jinja2 y CSS puro. Docker para despliegue "
     "reproducible.",
     PURPLE),
]):
    x = Inches(0.85 + i * 4.05)
    add_rect(s, x, col_y, col_w, col_h, BG_SOFT, line=BORDER)
    add_rect(s, x, col_y, Inches(0.12), col_h, color)
    add_text(s, x + Inches(0.3), col_y + Inches(0.2), col_w - Inches(0.4),
             Inches(0.5), title, size=16, bold=True, color=TEXT,
             font=FONT_TITLE)
    add_text(s, x + Inches(0.3), col_y + Inches(0.8), col_w - Inches(0.4),
             col_h - Inches(0.9), body, size=13, color=TEXT_DIM)

add_footer(s, 2)


# ═════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Temática
# ═════════════════════════════════════════════════════════════════════════
s = add_slide()
add_header(s, "02 · temática", "La corporación NEXUS Corp")

add_text(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(0.8),
         "Ambientación cyberpunk corporativa: una multinacional ficticia "
         "con portal de empleados, RRHH, servicios internos y un CEO "
         "Vault clasificado.",
         size=16, color=TEXT)

levels = [
    ("Level 1", "Portal de empleados", "Fuerza bruta (Hydra)", ACCENT),
    ("Level 2", "Base de datos RRHH", "SQL Injection (sqlmap)", RED),
    ("Level 3", "Hash del operador", "Hash cracking (Hashcat)", ORANGE),
    ("Level 4", "Servidor SSH interno", "Enumeración de archivos", GREEN),
    ("Level 5", "Directorio LDAP", "Enumeración LDAP", PURPLE),
    ("Level 6", "Vault API del CEO", "API fuzzing (ffuf)", ACCENT),
]
row_y = Inches(3.0)
row_h = Inches(0.55)
for i, (lvl, scenario, tool, color) in enumerate(levels):
    y = row_y + i * row_h
    add_rect(s, Inches(0.85), y, Inches(11.6), Inches(0.5), BG_SOFT,
             line=BORDER)
    add_rect(s, Inches(0.85), y, Inches(0.1), Inches(0.5), color)
    add_text(s, Inches(1.05), y + Inches(0.1), Inches(1.4), Inches(0.3),
             lvl, size=13, bold=True, color=color, font=FONT_TITLE)
    add_text(s, Inches(2.55), y + Inches(0.1), Inches(5.5), Inches(0.3),
             scenario, size=13, color=TEXT)
    add_text(s, Inches(8.1), y + Inches(0.1), Inches(4.5), Inches(0.3),
             tool, size=13, color=TEXT_DIM, font=FONT_MONO)

add_footer(s, 3)


# ═════════════════════════════════════════════════════════════════════════
# Helper: vulnerability slides (what is + how to exploit)
# ═════════════════════════════════════════════════════════════════════════
def vuln_what(page, eyebrow, title, tagline, bullets, color):
    s = add_slide()
    add_header(s, eyebrow, title, accent=color)
    add_text(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(0.8),
             tagline, size=17, color=TEXT)
    add_bullets(s, Inches(0.95), Inches(3.0), Inches(11.5), Inches(3.8),
                bullets, size=15, color=TEXT, bullet_color=color)
    add_footer(s, page)


def vuln_how(page, eyebrow, title, steps, code, note, color):
    s = add_slide()
    add_header(s, eyebrow, title, accent=color)

    # Left column: steps
    add_text(s, Inches(0.85), Inches(1.85), Inches(5.3), Inches(0.4),
             "PASOS", size=11, bold=True, color=color, font=FONT_TITLE)
    add_bullets(s, Inches(0.85), Inches(2.25), Inches(5.3), Inches(4.0),
                steps, size=14, color=TEXT, bullet_color=color)

    # Right column: command
    add_text(s, Inches(6.5), Inches(1.85), Inches(6.0), Inches(0.4),
             "COMANDO", size=11, bold=True, color=color, font=FONT_TITLE)
    add_code(s, Inches(6.5), Inches(2.25), Inches(6.0), Inches(3.3),
             code, size=11)

    if note:
        add_rect(s, Inches(6.5), Inches(5.7), Inches(6.0), Inches(1.15),
                 BG_SOFT, line=BORDER)
        add_rect(s, Inches(6.5), Inches(5.7), Inches(0.1),
                 Inches(1.15), color)
        add_text(s, Inches(6.7), Inches(5.75), Inches(5.8), Inches(0.3),
                 "PALABRA CLAVE EN DICCIONARIO", size=10, bold=True,
                 color=color, font=FONT_TITLE)
        add_text(s, Inches(6.7), Inches(6.05), Inches(5.8), Inches(0.8),
                 note, size=12, color=TEXT)

    add_footer(s, page)


# ═════════════════════════════════════════════════════════════════════════
# LEVEL 1 — Brute-force (slides 4, 5)
# ═════════════════════════════════════════════════════════════════════════
vuln_what(
    page=4,
    eyebrow="Level 1 · ¿qué es?",
    title="Fuerza bruta en autenticación",
    tagline="Un formulario de login sin límite de intentos ni captcha "
            "permite probar miles de combinaciones usuario/contraseña "
            "hasta encontrar una válida.",
    bullets=[
        "Raíz del problema: ausencia de rate-limiting, lockout y MFA.",
        "Contraseñas débiles o por defecto (admin/admin, root/toor) "
        "aceleran el ataque.",
        "OWASP Top 10 — A07:2021 Identification & Authentication Failures.",
        "Impacto: toma total de cuentas, escalado a datos internos.",
    ],
    color=ACCENT,
)

vuln_how(
    page=5,
    eyebrow="Level 1 · exploit",
    title="Hydra contra el portal de empleados",
    steps=[
        "Identificar el endpoint de login y los campos del formulario.",
        "Detectar el mensaje de fallo (\"Login failed\") para filtrar.",
        "Lanzar Hydra con usuario fijo y diccionario de contraseñas.",
        "Probar manualmente la credencial obtenida.",
    ],
    code=(
        "hydra -l admin \\\n"
        "  -P /usr/share/wordlists/rockyou.txt \\\n"
        "  TARGET http-post-form \\\n"
        "  \"/level/1/login:\"\\\n"
        "  \"username=^USER^&password=^PASS^:\"\\\n"
        "  \"Login failed\""
    ),
    note="La palabra \"admin\" debe estar presente en el diccionario. "
         "rockyou.txt ya la contiene.",
    color=ACCENT,
)


# ═════════════════════════════════════════════════════════════════════════
# LEVEL 2 — SQL Injection (slides 6, 7)
# ═════════════════════════════════════════════════════════════════════════
vuln_what(
    page=6,
    eyebrow="Level 2 · ¿qué es?",
    title="SQL Injection",
    tagline="La aplicación concatena input del usuario directamente en "
            "una consulta SQL, permitiendo al atacante modificar la "
            "lógica de la query y leer tablas arbitrarias.",
    bullets=[
        "Causa: consultas construidas con f-strings o concatenación en "
        "lugar de parámetros.",
        "Técnicas: UNION-based, boolean-blind, error-based, "
        "time-based.",
        "OWASP Top 10 — A03:2021 Injection.",
        "Impacto: exfiltración de bases de datos completas, escritura "
        "y RCE en algunos motores.",
    ],
    color=RED,
)

vuln_how(
    page=7,
    eyebrow="Level 2 · exploit",
    title="sqlmap contra /level/2/search",
    steps=[
        "Detectar el parámetro vulnerable (search) en modo JSON.",
        "sqlmap identifica UNION-based automáticamente.",
        "Enumerar tablas y volcar ssh_credentials.",
        "Obtener el hash MD5 de sysop para el Level 3.",
    ],
    code=(
        "sqlmap -u \\\n"
        "  \"http://TARGET:5000/level/2/search\"\\\n"
        "  \"?search=test&format=json\" \\\n"
        "  -T ssh_credentials \\\n"
        "  --dump --batch"
    ),
    note=None,
    color=RED,
)


# ═════════════════════════════════════════════════════════════════════════
# LEVEL 3 — Hash cracking (slides 8, 9)
# ═════════════════════════════════════════════════════════════════════════
vuln_what(
    page=8,
    eyebrow="Level 3 · ¿qué es?",
    title="Hash cracking de contraseñas",
    tagline="Cuando se filtra una base de datos con hashes de "
            "contraseñas, un atacante puede revertirlos offline si el "
            "algoritmo es débil o no usa salt.",
    bullets=[
        "MD5 y SHA1 sin salt se rompen a miles de millones de hashes "
        "por segundo con GPU.",
        "Diccionarios + reglas (rockyou + rules) cubren la mayoría de "
        "contraseñas humanas.",
        "OWASP Top 10 — A02:2021 Cryptographic Failures.",
        "Impacto: reutilización de credenciales en otros servicios "
        "(credential stuffing).",
    ],
    color=ORANGE,
)

vuln_how(
    page=9,
    eyebrow="Level 3 · exploit",
    title="Hashcat contra MD5 de sysop",
    steps=[
        "Guardar el hash obtenido del Level 2 en un archivo.",
        "Lanzar hashcat en modo 0 (MD5) contra el diccionario.",
        "Si rockyou.txt no lo contiene, añadir la contraseña a un "
        "wordlist custom.",
        "Enviar la contraseña encontrada en /level/3/verify.",
    ],
    code=(
        "echo \"98ae336a33cb54a3d5effde7\"\\\n"
        "     \"f32c06c8\" > hash.txt\n"
        "hashcat -m 0 hash.txt \\\n"
        "  /usr/share/wordlists/rockyou.txt\n\n"
        "# Si no está en rockyou:\n"
        "echo \"Op3r4t0r!\" >> custom.txt\n"
        "hashcat -m 0 hash.txt custom.txt"
    ),
    note="La palabra exacta a encontrar es Op3r4t0r! — no existe en "
         "rockyou.txt, requiere un wordlist custom.",
    color=ORANGE,
)


# ═════════════════════════════════════════════════════════════════════════
# LEVEL 4 — SSH / file enumeration (slides 10, 11)
# ═════════════════════════════════════════════════════════════════════════
vuln_what(
    page=10,
    eyebrow="Level 4 · ¿qué es?",
    title="Enumeración post-explotación en SSH",
    tagline="Una vez dentro del host, archivos de configuración, "
            "historiales y dumps de debug dejados por los "
            "desarrolladores revelan credenciales y arquitectura "
            "interna.",
    bullets=[
        "Problema: información sensible en el filesystem "
        "(.bash_history, /tmp, notes.txt, backup.key).",
        "Permisos laxos exponen estos archivos a cualquier usuario "
        "autenticado.",
        "OWASP Top 10 — A05:2021 Security Misconfiguration.",
        "Impacto: pivoting hacia otros servicios (LDAP, DBs, APIs "
        "internas).",
    ],
    color=GREEN,
)

vuln_how(
    page=11,
    eyebrow="Level 4 · exploit",
    title="Exploración del filesystem",
    steps=[
        "Login simulado con sysop / Op3r4t0r! (obtenido en Level 3).",
        "Listar el home y leer notes.txt y .bash_history.",
        "Consultar /etc/nexus/ldap.conf y backup.key para el Level 5.",
        "Un único comando devuelve la flag del reto.",
    ],
    code=(
        "$ ls\n"
        "$ cat notes.txt\n"
        "$ cat .bash_history\n"
        "$ cat /etc/nexus/ldap.conf\n"
        "$ cat /etc/nexus/backup.key\n"
        "\n"
        "# Comando que entrega la flag:\n"
        "$ cat /tmp/debug_dump.txt"
    ),
    note=None,
    color=GREEN,
)


# ═════════════════════════════════════════════════════════════════════════
# LEVEL 5 — LDAP enumeration (slides 12, 13)
# ═════════════════════════════════════════════════════════════════════════
vuln_what(
    page=12,
    eyebrow="Level 5 · ¿qué es?",
    title="Enumeración LDAP",
    tagline="Un directorio LDAP mal configurado (anonymous bind, ACLs "
            "amplias) permite listar usuarios, grupos y hasta hashes "
            "de contraseñas.",
    bullets=[
        "Problema: anonymous bind habilitado y credenciales "
        "administrativas reutilizadas.",
        "Con credenciales de admin LDAP se exponen atributos "
        "sensibles (userPassword).",
        "OWASP Top 10 — A01:2021 Broken Access Control.",
        "Impacto: acceso a cuentas de servicio con privilegios "
        "elevados.",
    ],
    color=PURPLE,
)

vuln_how(
    page=13,
    eyebrow="Level 5 · exploit",
    title="Bind administrativo contra el directorio",
    steps=[
        "Primer query anónimo revela la estructura pero oculta hashes.",
        "Repetir con credenciales admin obtenidas en Level 4.",
        "Localizar la entrada vault_svc (clearance_level = 5).",
        "Copiar el hash SHA256 de userPassword para el Level 6.",
    ],
    code=(
        "# Parámetros del formulario /level/5\n"
        "Bind DN:   cn=admin,\n"
        "           dc=nexuscorp,\n"
        "           dc=local\n"
        "Password:  Ldap@dm1n_2024\n"
        "Base DN:   dc=nexuscorp,dc=local\n"
        "Filter:    (objectClass=*)\n"
        "Attrs:     *"
    ),
    note=None,
    color=PURPLE,
)


# ═════════════════════════════════════════════════════════════════════════
# LEVEL 6 — API Fuzzing (slides 14, 15)
# ═════════════════════════════════════════════════════════════════════════
vuln_what(
    page=14,
    eyebrow="Level 6 · ¿qué es?",
    title="API fuzzing para descubrir recursos ocultos",
    tagline="Endpoints y parámetros no documentados que la aplicación "
            "sirve si se conocen sus nombres. Un atacante los "
            "descubre enviando miles de peticiones con un diccionario.",
    bullets=[
        "Security through obscurity: la API existe pero no está "
        "listada públicamente.",
        "Autenticación única por header (X-Vault-Token) sin rotación.",
        "OWASP API Top 10 — API9:2023 Improper Inventory Management.",
        "Impacto: acceso a documentos clasificados, logs, backups.",
    ],
    color=ACCENT,
)

vuln_how(
    page=15,
    eyebrow="Level 6 · exploit",
    title="ffuf en dos fases: endpoints + doc_id",
    steps=[
        "Crackear el hash SHA256 de vault_svc (V4ult_M4st3r_K3y!).",
        "Descargar el wordlist público de la API.",
        "Fase 1 — descubrir endpoints válidos filtrando 404.",
        "Fase 2 — fuzzear doc_id (1–95) filtrando por tamaño común.",
        "doc_id=73 devuelve los documentos clasificados con la flag.",
    ],
    code=(
        "TOKEN='V4ult_M4st3r_K3y!'\n"
        "\n"
        "ffuf -w api-wordlist.txt \\\n"
        "  -H \"X-Vault-Token: $TOKEN\" \\\n"
        "  -u TARGET/level/6/vault-api/FUZZ \\\n"
        "  -fc 404\n"
        "\n"
        "seq 1 95 > nums.txt\n"
        "ffuf -w nums.txt \\\n"
        "  -H \"X-Vault-Token: $TOKEN\" \\\n"
        "  -u TARGET/.../archives?doc_id=FUZZ \\\n"
        "  -fs 95"
    ),
    note="La palabra exacta a encontrar (previo hash crack) es "
         "V4ult_M4st3r_K3y! — requiere wordlist custom. "
         "El doc_id ganador es 73.",
    color=ACCENT,
)


# ═════════════════════════════════════════════════════════════════════════
# SLIDE 16 — ¿Por qué es importante asegurarlas?
# ═════════════════════════════════════════════════════════════════════════
s = add_slide()
add_header(s, "03 · impacto", "¿Por qué asegurar cada una?")

reasons = [
    ("Brute-force",
     "Compromete cuentas administrativas en minutos. Pivote a todo "
     "el stack.",
     ACCENT),
    ("SQL Injection",
     "Exfiltración total de la base de datos. Cumplimiento GDPR, "
     "LGPD, HIPAA.",
     RED),
    ("Hashes débiles",
     "Credential stuffing contra otros servicios (Office 365, "
     "bancos, VPN).",
     ORANGE),
    ("Misconfig / archivos",
     "Filtra arquitectura y credenciales internas; acelera el "
     "movimiento lateral.",
     GREEN),
    ("LDAP abierto",
     "El directorio es la columna vertebral del SSO corporativo. "
     "Compromiso = dominio.",
     PURPLE),
    ("APIs ocultas",
     "Shadow endpoints bypassan WAF y logging; exponen datos "
     "confidenciales sin auditar.",
     ACCENT),
]

cols = 3
col_w = Inches(3.95)
row_h = Inches(1.55)
for i, (name, desc, color) in enumerate(reasons):
    col = i % cols
    row = i // cols
    x = Inches(0.85 + col * 4.05)
    y = Inches(2.0 + row * 1.7)
    add_rect(s, x, y, col_w, row_h, BG_SOFT, line=BORDER)
    add_rect(s, x, y, Inches(0.1), row_h, color)
    add_text(s, x + Inches(0.3), y + Inches(0.15),
             col_w - Inches(0.4), Inches(0.4),
             name, size=14, bold=True, color=color, font=FONT_TITLE)
    add_text(s, x + Inches(0.3), y + Inches(0.55),
             col_w - Inches(0.4), row_h - Inches(0.6),
             desc, size=12, color=TEXT)

add_text(s, Inches(0.85), Inches(5.65), Inches(11.6), Inches(1.2),
         "Cada vulnerabilidad es una puerta. Una sola puerta abierta "
         "permite al atacante encadenar las demás — exactamente lo que "
         "simula este CTF.",
         size=15, color=TEXT_DIM)

add_footer(s, 16)


# ═════════════════════════════════════════════════════════════════════════
# SLIDE 17 — ¿Cómo asegurarlas?
# ═════════════════════════════════════════════════════════════════════════
s = add_slide()
add_header(s, "04 · mitigación", "La forma más sencilla de asegurarlas")

# Top row: libraries
add_text(s, Inches(0.85), Inches(1.8), Inches(11.6), Inches(0.4),
         "BIBLIOTECAS ESPECÍFICAS (Python / Flask)",
         size=11, bold=True, color=ACCENT, font=FONT_TITLE)

libs = [
    ("Flask-Limiter",
     "Rate-limiting por IP y por usuario. Cierra Level 1 y 6.",
     ACCENT),
    ("SQLAlchemy ORM",
     "Queries parametrizadas por defecto. Cierra Level 2.",
     RED),
    ("Argon2-cffi / bcrypt",
     "Hash lento con salt. Cierra Level 3.",
     ORANGE),
    ("Flask-Talisman + "
     "python-dotenv",
     "Headers seguros y secretos fuera del repo. Cierra Level 4.",
     GREEN),
    ("ldap3 + ACLs "
     "estrictas",
     "Deshabilitar anonymous bind, mínimo privilegio. Cierra Level 5.",
     PURPLE),
    ("Flask-Smorest / "
     "APISpec",
     "Inventario y OpenAPI firmado. Cierra Level 6.",
     ACCENT),
]

cols = 3
col_w = Inches(3.95)
row_h = Inches(1.25)
for i, (name, desc, color) in enumerate(libs):
    col = i % cols
    row = i // cols
    x = Inches(0.85 + col * 4.05)
    y = Inches(2.25 + row * 1.35)
    add_rect(s, x, y, col_w, row_h, BG_SOFT, line=BORDER)
    add_rect(s, x, y, Inches(0.1), row_h, color)
    add_text(s, x + Inches(0.3), y + Inches(0.12),
             col_w - Inches(0.4), Inches(0.35),
             name, size=12, bold=True, color=color, font=FONT_MONO)
    add_text(s, x + Inches(0.3), y + Inches(0.5),
             col_w - Inches(0.4), row_h - Inches(0.55),
             desc, size=11, color=TEXT)

# Mindset box
add_rect(s, Inches(0.85), Inches(5.1), Inches(11.6), Inches(1.75),
         BG_SOFT, line=BORDER)
add_rect(s, Inches(0.85), Inches(5.1), Inches(0.12), Inches(1.75), ACCENT)
add_text(s, Inches(1.1), Inches(5.25), Inches(11.2), Inches(0.4),
         "MENTE CIBERSEGURA — desarrollo seguro por defecto",
         size=12, bold=True, color=ACCENT, font=FONT_TITLE)

mind_items = [
    "Asumir compromiso: validar entrada, negar por defecto, defensa "
    "en profundidad.",
    "Principio de mínimo privilegio en credenciales, roles y rutas.",
    "Secrets fuera del código; rotación automática; MFA en todo acceso "
    "administrativo.",
    "Logging + alertas en endpoints sensibles; revisar SIEM semanalmente.",
    "SAST + DAST + dependencias auditadas (pip-audit, Bandit) en CI/CD.",
]
tb = s.shapes.add_textbox(Inches(1.1), Inches(5.6),
                          Inches(11.2), Inches(1.2))
tf = tb.text_frame
tf.word_wrap = True
for i, item in enumerate(mind_items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.line_spacing = 1.2
    rb = p.add_run()
    rb.text = "▸  "
    rb.font.name = FONT_BODY
    rb.font.size = Pt(12)
    rb.font.color.rgb = ACCENT
    rb.font.bold = True
    r = p.add_run()
    r.text = item
    r.font.name = FONT_BODY
    r.font.size = Pt(12)
    r.font.color.rgb = TEXT

add_footer(s, 17)


# ═════════════════════════════════════════════════════════════════════════
# Save
# ═════════════════════════════════════════════════════════════════════════
out = Path(__file__).resolve().parent.parent / "NexusCorp_CTF_Presentation.pptx"
prs.save(out)
print(f"Presentation saved: {out}")
print(f"Total slides: {len(prs.slides)}")
